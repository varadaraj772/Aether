import os
import sys
import tempfile
import subprocess
import math
import hashlib
import time

import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

from models import DOCUMENT_PROMPT, MODEL_DIMENSION, MODEL_NAME, get_model
from lance_store import db, db_lock

AETHER_DIR = os.environ.get("AETHER_DIR", os.path.join(os.path.expanduser("~"), ".aether"))
os.makedirs(os.path.join(AETHER_DIR, "db"), exist_ok=True)

TABLE_NAME = "documents"

_profile_enabled = os.environ.get("AETHER_PROFILE_INDEX") == "1"
_profile = {
    "files": 0,
    "chunks": 0,
    "extract_seconds": 0.0,
    "hash_seconds": 0.0,
    "embedding_seconds": 0.0,
    "validation_seconds": 0.0,
    "database_seconds": 0.0,
    "embedding_calls": 0,
    "embedding_texts": 0,
}


def reset_index_profile():
    for key in _profile:
        _profile[key] = 0 if key in {"files", "chunks", "embedding_calls", "embedding_texts"} else 0.0


def get_index_profile():
    return dict(_profile)


def _profile_add(key, elapsed):
    if _profile_enabled:
        _profile[key] += elapsed

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".log",
    ".swift",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".rs",
    ".go",
    ".java",
    ".kt",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".m",
    ".mm",
    ".rb",
    ".php",
    ".sh",
    ".yaml",
    ".yml",
    ".toml",
    ".xml",
    ".sql",
    ".r",
    ".scala",
    ".vue",
    ".svelte",
    ".eml",
}

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pages",
    ".xlsx",
    ".pptx",
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
    ".mp3",
    ".wav",
    ".flac",
    ".m4a",
    ".aac",
    ".ogg",
    ".opus",
    ".mp4",
    ".mov",
    ".m4v",
    ".avi",
    ".mkv",
    ".webm",
    *TEXT_EXTENSIONS,
}

PDF_OCR_MIN_CHARS = 40
PDF_OCR_MAX_PAGES = 10
DEFAULT_MAX_TEXT_BYTES = 10 * 1024 * 1024

EXCLUDED_DIR_NAMES = {
    ".git",
    ".svn",
    ".hg",
    "node_modules",
    "__pycache__",
    ".venv",
    "venv",
    "dist",
    "build",
    ".cache",
    ".Trash",
    ".idea",
    ".cursor",
    "site-packages",
}


def _active_excluded_names():
    try:
        from folders import load_excluded_dir_names

        names = load_excluded_dir_names()
        if names:
            return names
    except Exception:
        pass
    return set(EXCLUDED_DIR_NAMES)


def _should_skip_dir(dirname):
    excluded = _active_excluded_names()
    return any(dirname == name or dirname.lower() == name.lower() for name in excluded)


def should_skip_path(path):
    excluded = _active_excluded_names()
    normalized = os.path.normpath(path)
    parts = normalized.split(os.sep)
    for name in excluded:
        needle = name.strip().strip("/\\")
        if not needle:
            continue
        if any(part == needle or part.lower() == needle.lower() for part in parts):
            return True
        lower_path = normalized.lower()
        lower_needle = needle.lower()
        if lower_path.endswith("/" + lower_needle) or lower_path.endswith("\\" + lower_needle):
            return True
        if lower_path == lower_needle:
            return True
    return False



def extract_pages_text(path):
    temp_dir = tempfile.mkdtemp(prefix="aether-pages-")
    try:
        pdf_path = os.path.join(temp_dir, os.path.basename(path) + ".pdf")

        # Pass paths as AppleScript string literals, never by interpolating
        # untrusted filenames into an expression.
        def applescript_string(value):
            return '"' + value.replace('\\', '\\\\').replace('"', '\\"') + '"'
        source = applescript_string(path)
        destination = applescript_string(pdf_path)
        script = f'''
        tell application "Pages"
            open POSIX file {source}
            export front document to POSIX file {destination} as PDF
            close front document saving no
        end tell
        '''

        subprocess.run(
            ["osascript", "-e", script],
            check=True,
            capture_output=True,
            text=True
        )

        reader = PdfReader(pdf_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        print(f"\nPAGES ERROR: {path}")
        print(e)
        return ""
    finally:
        import shutil
        shutil.rmtree(temp_dir, ignore_errors=True)


def extract_pdf(path, max_pages=None, ocr_max_pages=PDF_OCR_MAX_PAGES):
    text = ""
    try:
        reader = PdfReader(path)
        pages = reader.pages if max_pages is None else reader.pages[:max_pages]
        text = "\n".join(page.extract_text() or "" for page in pages)
    except Exception as e:
        print(f"\nPDF ERROR: {path}")
        print(e)
        text = ""

    if len(text.strip()) < PDF_OCR_MIN_CHARS:
        ocr_text = ocr_pdf_pages(path, max_pages=ocr_max_pages)
        if len(ocr_text.strip()) > len(text.strip()):
            return ocr_text

    return text


def extract_docx(path):
    try:
        doc = Document(path)
        parts = []
        # Body order is preserved for paragraphs and tables as represented by
        # python-docx's XML tree, rather than silently dropping tables.
        for child in doc.element.body.iterchildren():
            if child.tag.endswith('}p'):
                parts.append("".join(node.text or "" for node in child.iter() if node.tag.endswith('}t')))
            elif child.tag.endswith('}tbl'):
                for row in child.iterchildren():
                    cells = []
                    for cell in row.iterchildren():
                        cells.append(" ".join(node.text or "" for node in cell.iter() if node.tag.endswith('}t')))
                    if cells:
                        parts.append("\t".join(cells))
        return "\n".join(part for part in parts if part.strip())
    except Exception as e:
        print(f"\nDOCX ERROR: {path}")
        print(e)
        return ""


def extract_text_file(path):
    try:
        max_bytes = int(os.environ.get("AETHER_MAX_TEXT_BYTES", str(DEFAULT_MAX_TEXT_BYTES)))
        if os.path.getsize(path) > max_bytes:
            raise ExtractionError(f"Text file exceeds configured limit ({max_bytes} bytes): {path}")
        with open(path, "rb") as f:
            sample = f.read(8192)
            if b"\x00" in sample:
                raise ExtractionError(f"Binary content is not searchable text: {path}")
            return (sample + f.read()).decode("utf-8", errors="replace")
    except ExtractionError:
        raise
    except Exception as e:
        print(f"\nTEXT ERROR: {path}")
        print(e)
        return ""


def extract_xlsx(path):
    try:
        sheets = pd.read_excel(path, sheet_name=None, header=None)
        parts = []
        for name, frame in sheets.items():
            values = frame.fillna("").astype(str).to_string(index=False, header=False)
            parts.append(f"[Sheet: {name}]\n{values}" if values.strip() else f"[Sheet: {name}]")
        return "\n\n".join(parts)
    except Exception as e:
        print(f"\nXLSX ERROR: {path}")
        print(e)
        return ""


def extract_pptx(path):
    try:
        prs = Presentation(path)
        text = ""
        for slide_number, slide in enumerate(prs.slides, 1):
            slide_parts = [f"[Slide {slide_number}]"]
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    if shape.text.strip():
                        slide_parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_parts.append("\t".join(cell.text for cell in row.cells))
            text += "\n".join(slide_parts) + "\n"
        return text
    except Exception as e:
        print(f"\nPPTX ERROR: {path}")
        print(e)
        return ""


_ocr_engine = None

def extract_image(path):
    global _ocr_engine
    try:
        # pyrefly: ignore [missing-import]
        from rapidocr_onnxruntime import RapidOCR
        if _ocr_engine is None:
            _ocr_engine = RapidOCR()
        result, _ = _ocr_engine(path)
        if result:
            return "\n".join(line[1] for line in result if line and len(line) > 1)
    except Exception as e:
        print(f"\nOCR ERROR: {path}")
        print(e)
    return ""


def ocr_pdf_pages(path, max_pages=PDF_OCR_MAX_PAGES):
    """Render PDF pages to images and OCR them. Prefers pypdfium2."""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        print(f"\nPDF OCR SKIP (pypdfium2 missing): {path}", file=sys.stderr)
        return ""

    texts = []
    pdf = None
    try:
        pdf = pdfium.PdfDocument(path)
        page_count = min(len(pdf), max_pages)
        for i in range(page_count):
            page = pdf[i]
            try:
                bitmap = page.render(scale=2.0)
                image = bitmap.to_pil()
            finally:
                try:
                    page.close()
                except Exception:
                    pass

            fd, tmp_path = tempfile.mkstemp(suffix=".png")
            os.close(fd)
            try:
                image.save(tmp_path, format="PNG")
                page_text = extract_image(tmp_path)
                if page_text.strip():
                    texts.append(page_text)
            finally:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass
    except Exception as e:
        print(f"\nPDF OCR ERROR: {path}")
        print(e)
        return ""
    finally:
        if pdf is not None:
            try:
                pdf.close()
            except Exception:
                pass

    return "\n".join(texts)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return extract_pdf(path)
    elif ext == ".docx":
        return extract_docx(path)
    elif ext == ".pages":
        return extract_pages_text(path)
    elif ext == ".xlsx":
        return extract_xlsx(path)
    elif ext == ".pptx":
        return extract_pptx(path)
    elif ext in [".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"]:
        return extract_image(path)
    elif ext in TEXT_EXTENSIONS:
        return extract_text_file(path)

    return ""


class ExtractionError(RuntimeError):
    """The source could not be read reliably enough to replace its index."""


def chunk_text(text, chunk_size=800, overlap=150):
    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += (chunk_size - overlap)

    return chunks


def _normalize_folder(folder):
    return os.path.normpath(folder.rstrip(os.sep)) or folder


def _escape_sql(value):
    return value.replace("'", "''")


def _file_sha256(path, block_size=1024 * 1024):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        while block := handle.read(block_size):
            digest.update(block)
    return digest.hexdigest()


def _projected_rows(table, *, where=None, columns=None, limit=None):
    """Read only requested columns, using LanceDB's filtered query when available."""
    try:
        query = table.search()
        if where:
            query = query.where(where)
        if columns:
            query = query.select(columns)
        if limit is not None:
            query = query.limit(limit)
        return query.to_list()
    except Exception:
        try:
            query = table.query()
            if where:
                query = query.where(where)
            if columns:
                query = query.select(columns)
            if limit is not None:
                query = query.limit(limit)
            return query.to_list()
        except Exception:
            # Compatibility fallback for older LanceDB/table fakes. The
            # projection is still applied before returning to callers.
            rows = table.to_arrow().to_pylist()
            return [
                {key: row.get(key) for key in (columns or row.keys())}
                for row in rows
            ]


def _path_in_folder(path, folder):
    folder = _normalize_folder(folder)
    path = os.path.normpath(path)
    return path == folder or path.startswith(folder + os.sep)


def delete_folder_from_index(folder):
    folder = _normalize_folder(folder)
    prefix = _escape_sql(folder + os.sep)
    escaped = _escape_sql(folder)

    with db_lock:
        if TABLE_NAME not in db.table_names():
            return
        table = db.open_table(TABLE_NAME)
        try:
            table.delete(f"path LIKE '{prefix}%' OR path = '{escaped}'")
            return
        except Exception:
            pass

        try:
            paths = [row.get("path") for row in _projected_rows(table, columns=["path"])]
            for path in {p for p in paths if _path_in_folder(p, folder)}:
                table.delete(f"path = '{_escape_sql(path)}'")
        except Exception as e:
            print(f"DELETE FOLDER ERROR: {folder}")
            print(e)


def folder_is_indexed(folder):
    folder = _normalize_folder(folder)
    try:
        with db_lock:
            if TABLE_NAME not in db.table_names():
                return False
            table = db.open_table(TABLE_NAME)
            paths = [row.get("path") for row in _projected_rows(table, columns=["path"])]
        return any(_path_in_folder(path, folder) for path in paths)
    except Exception:
        return False


def delete_file_from_index(path):
    path = os.path.normpath(path)
    try:
        with db_lock:
            if TABLE_NAME not in db.table_names():
                return
            table = db.open_table(TABLE_NAME)
            table.delete(f"path = '{_escape_sql(path)}'")
        print(f"Removed from index: {path}")
    except Exception as e:
        print(f"DELETE FILE ERROR: {path}")
        print(e)
        return

    folder = _parent_indexed_folder(path)
    if folder:
        from folders import get_folder_file_mtimes, set_folder_file_mtimes

        mtimes = get_folder_file_mtimes(folder)
        if mtimes.pop(path, None) is not None:
            set_folder_file_mtimes(folder, mtimes)


def get_folder_stats(folder):
    folder = _normalize_folder(folder)
    last_indexed = None
    try:
        from folders import get_folder_meta

        last_indexed = get_folder_meta(folder).get("last_indexed")
    except Exception:
        pass
    try:
        with db_lock:
            if TABLE_NAME not in db.table_names():
                return {
                    "indexed": False,
                    "file_count": 0,
                    "chunk_count": 0,
                    "last_indexed": last_indexed,
                }
            table = db.open_table(TABLE_NAME)
            paths = [row.get("path") for row in _projected_rows(table, columns=["path"])]
        folder_paths = [p for p in paths if _path_in_folder(p, folder)]
        return {
            "indexed": bool(folder_paths),
            "file_count": len(set(folder_paths)),
            "chunk_count": len(folder_paths),
            "last_indexed": last_indexed,
        }
    except Exception:
        return {
            "indexed": False,
            "file_count": 0,
            "chunk_count": 0,
            "last_indexed": last_indexed,
        }


def _walk_supported_files(folder):
    folder = _normalize_folder(folder)
    paths = []
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if not _should_skip_dir(d)]
        for file in files:
            path = os.path.join(root, file)
            if should_skip_path(path):
                continue
            ext = os.path.splitext(path)[1].lower()
            if ext in SUPPORTED_EXTENSIONS:
                paths.append(os.path.normpath(path))
    return paths


def _indexed_paths_in_folder(folder):
    folder = _normalize_folder(folder)
    try:
        with db_lock:
            if TABLE_NAME not in db.table_names():
                return set()
            table = db.open_table(TABLE_NAME)
            escaped = _escape_sql(folder + os.sep)
            exact = _escape_sql(folder)
            paths = [
                row.get("path")
                for row in _projected_rows(
                    table,
                    where=f"path LIKE '{escaped}%' OR path = '{exact}'",
                    columns=["path"],
                )
            ]
        return {os.path.normpath(p) for p in paths if _path_in_folder(p, folder)}
    except Exception:
        return set()


def _parent_indexed_folder(path):
    from folders import load_folders

    path = os.path.normpath(path)
    for folder in load_folders():
        folder = _normalize_folder(folder)
        if _path_in_folder(path, folder):
            return folder
    return None


def _write_index_rows(path, rows):
    """Replace one file's rows, restoring its prior rows if insertion fails."""
    started = time.perf_counter()
    with db_lock:
        if TABLE_NAME not in db.table_names():
            db.create_table(TABLE_NAME, data=rows)
        else:
            table = db.open_table(TABLE_NAME)
            try:
                schema_names = set(table.schema.names)
                rows_to_write = [{key: value for key, value in row.items() if key in schema_names} for row in rows]
            except Exception:
                rows_to_write = rows
            old_rows = _projected_rows(table, where=f"path = '{_escape_sql(path)}'")
            if old_rows:
                try:
                    table.delete(f"path = '{_escape_sql(path)}'")
                except Exception as exc:
                    raise RuntimeError(f"Could not remove old index rows for {path}") from exc
            try:
                table.add(rows_to_write)
            except Exception:
                if old_rows:
                    try:
                        table.add(old_rows)
                    except Exception as restore_error:
                        raise RuntimeError(
                            f"Index write and restoration failed for {path}: {restore_error}"
                        ) from restore_error
                raise
    _profile_add("database_seconds", time.perf_counter() - started)
    if _profile_enabled:
        _profile["files"] += 1
        _profile["chunks"] += len(rows)


def _index_file(path, model=None, text_vectors=None):
    path = os.path.normpath(path)
    if should_skip_path(path):
        return 0

    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return 0

    print(f"Processing: {path}")
    if not os.path.isfile(path):
        raise ExtractionError(f"File disappeared during indexing: {path}")
    started = time.perf_counter()
    try:
        text = extract_text(path)
    except Exception as exc:
        raise ExtractionError(f"Extraction failed for {path}: {exc}") from exc
    _profile_add("extract_seconds", time.perf_counter() - started)
    image_exts = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
    audio_exts = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".opus"}
    video_exts = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
    modality = None
    if ext in image_exts:
        modality = "image"
    elif ext in audio_exts:
        modality = "audio"
    elif ext in video_exts:
        modality = "video"

    # Media can be semantically indexed even when it has no OCR/transcript.
    if not text.strip() and modality is None:
        # An empty extraction is ambiguous (empty file, unreadable file, or a
        # parser failure). Preserve the last known-good rows and let the
        # caller report this file as failed.
        raise ExtractionError(f"No searchable content extracted from {path}")

    if model is None:
        model = get_model()

    started = time.perf_counter()
    try:
        stat = os.stat(path)
        fingerprint = _file_sha256(path)
    except OSError as exc:
        raise ExtractionError(f"File changed or disappeared during indexing: {path}") from exc
    _profile_add("hash_seconds", time.perf_counter() - started)

    chunks = chunk_text(text) if text.strip() else [""]
    rows = []
    if modality is None and text_vectors is None:
        batch_size = max(1, int(os.environ.get("AETHER_EMBED_BATCH_SIZE", "16")))
        started = time.perf_counter()
        text_vectors = _encode_text_batches(model, chunks, batch_size)
        _profile_add("embedding_seconds", time.perf_counter() - started)
    else:
        text_vectors = []
    for chunk_id, chunk in enumerate(chunks):
        if modality:
            media_kwargs = {"show_progress_bar": False}
            if modality == "video":
                media_kwargs["processing_kwargs"] = {
                    "video": {
                        "max_frames": 32,
                        "fps": 1.0,
                        "overflow_strategy": "uniform",
                        "add_timestamps": True,
                    }
                }
            vector = model.encode(
                {modality: path},
                **media_kwargs,
            ).tolist()
        else:
            vector = text_vectors[chunk_id]
        rows.append({
            "path": path,
            "file_name": os.path.basename(path),
            "chunk_id": chunk_id,
            "content": chunk,
            "modality": modality or "text",
            "vector": vector,
            "file_size": stat.st_size,
            "mtime": stat.st_mtime,
            "fingerprint": fingerprint,
            "embedding_model": MODEL_NAME,
            "embedding_dimension": MODEL_DIMENSION,
            "extraction_status": "complete",
        })

    started = time.perf_counter()
    _validate_index_rows(rows, path)
    _profile_add("validation_seconds", time.perf_counter() - started)

    # Prepare and validate outside the lock. Keep the old rows in memory while
    # replacing them so a failed write can restore the previous document.
    _write_index_rows(path, rows)
    print(f"Indexed: {path}")
    return len(rows)


def _validate_index_rows(rows, path):
    """Reject malformed rows before touching LanceDB."""
    if not rows:
        raise ValueError(f"No index rows prepared for {path}")
    for expected_chunk_id, row in enumerate(rows):
        vector = row.get("vector")
        if row.get("path") != path or row.get("chunk_id") != expected_chunk_id:
            raise ValueError(f"Invalid index metadata prepared for {path}")
        if not isinstance(vector, list) or len(vector) != MODEL_DIMENSION:
            raise ValueError(f"Invalid embedding dimension for {path}")
        if not all(isinstance(value, (int, float)) and math.isfinite(value) for value in vector):
            raise ValueError(f"Invalid embedding values for {path}")


def _encode_text_batches(model, chunks, batch_size):
    """Encode bounded batches, falling back to one item if lists are unsupported."""
    from index_progress import is_cancelled

    vectors = []
    for start in range(0, len(chunks), batch_size):
        if is_cancelled():
            raise InterruptedError("Indexing cancelled")
        batch = chunks[start : start + batch_size]
        kwargs = {"prompt_name": DOCUMENT_PROMPT, "show_progress_bar": False}
        try:
            encoded = model.encode(batch, **kwargs)
            encoded = encoded.tolist() if hasattr(encoded, "tolist") else encoded
            if len(encoded) != len(batch):
                raise ValueError("Embedding batch length did not match chunk count")
            vectors.extend(encoded)
            if _profile_enabled:
                _profile["embedding_calls"] += 1
                _profile["embedding_texts"] += len(batch)
        except (TypeError, ValueError, AttributeError):
            for chunk in batch:
                if is_cancelled():
                    raise InterruptedError("Indexing cancelled")
                encoded = model.encode(chunk, **kwargs)
                if _profile_enabled:
                    _profile["embedding_calls"] += 1
                    _profile["embedding_texts"] += 1
                vectors.append(encoded.tolist() if hasattr(encoded, "tolist") else encoded)
    return vectors


def _index_text_batch(paths, model, batch_size):
    """Index text files using bounded batches shared across file boundaries."""
    from index_progress import is_cancelled

    prepared = []
    texts = []
    for path in paths:
        if is_cancelled():
            raise InterruptedError("Indexing cancelled")
        path = os.path.normpath(path)
        if not os.path.isfile(path):
            raise ExtractionError(f"File disappeared during indexing: {path}")
        started = time.perf_counter()
        text = extract_text(path)
        _profile_add("extract_seconds", time.perf_counter() - started)
        if not text.strip():
            raise ExtractionError(f"No searchable content extracted from {path}")
        started = time.perf_counter()
        try:
            stat = os.stat(path)
            fingerprint = _file_sha256(path)
        except OSError as exc:
            raise ExtractionError(f"File changed or disappeared during indexing: {path}") from exc
        _profile_add("hash_seconds", time.perf_counter() - started)
        chunks = chunk_text(text)
        prepared.append((path, chunks, stat, fingerprint))
        texts.extend(chunks)

    started = time.perf_counter()
    vectors = _encode_text_batches(model, texts, batch_size)
    _profile_add("embedding_seconds", time.perf_counter() - started)
    if is_cancelled():
        raise InterruptedError("Indexing cancelled")
    offset = 0
    counts = {}
    write_error = None
    for path, chunks, stat, fingerprint in prepared:
        file_vectors = vectors[offset : offset + len(chunks)]
        offset += len(chunks)
        rows = []
        for chunk_id, (chunk, vector) in enumerate(zip(chunks, file_vectors)):
            rows.append({
                "path": path,
                "file_name": os.path.basename(path),
                "chunk_id": chunk_id,
                "content": chunk,
                "modality": "text",
                "vector": vector,
                "file_size": stat.st_size,
                "mtime": stat.st_mtime,
                "fingerprint": fingerprint,
                "embedding_model": MODEL_NAME,
                "embedding_dimension": MODEL_DIMENSION,
                "extraction_status": "complete",
            })
        started = time.perf_counter()
        _validate_index_rows(rows, path)
        _profile_add("validation_seconds", time.perf_counter() - started)
        try:
            _write_index_rows(path, rows)
            counts[path] = len(rows)
        except InterruptedError:
            raise
        except Exception as exc:
            # LanceDB does not provide a transaction spanning documents. Keep
            # processing the remaining validated documents, then surface the
            # first failure so the caller can report/retry that document.
            if write_error is None:
                write_error = exc
    if write_error is not None:
        raise write_error
    return counts


def _build_index(folder, force=False):
    folder = _normalize_folder(folder)
    from folders import get_folder_file_mtimes, set_folder_file_mtimes, update_folder_meta
    from index_progress import is_cancelled, set_progress

    disk_paths = sorted(_walk_supported_files(folder))
    disk_path_set = set(disk_paths)
    model = get_model()

    # EmbeddingGemma 2 changes the vector width from MiniLM's 384 dimensions.
    # Drop the old table once, then let this normal index pass rebuild it.
    with db_lock:
        if TABLE_NAME in db.table_names():
            try:
                sample = db.open_table(TABLE_NAME).to_arrow().column("vector")[0].as_py()
                if len(sample) != MODEL_DIMENSION:
                    db.drop_table(TABLE_NAME)
                    force = True
            except Exception:
                # An unreadable or legacy schema is not safe to append to.
                # Drop it before the first replacement so the new model space
                # gets a correctly typed table instead of failing per file.
                try:
                    if TABLE_NAME in db.table_names():
                        db.drop_table(TABLE_NAME)
                except Exception:
                    pass
                force = True

    indexed_paths = _indexed_paths_in_folder(folder)
    meta_mtimes = get_folder_file_mtimes(folder)

    total_steps = len(indexed_paths - disk_path_set) + len(disk_paths)
    set_progress(active=True, current=0, total=max(total_steps, 1), folder=folder, file="")
    step = 0

    removed = 0
    for path in indexed_paths - disk_path_set:
        if is_cancelled():
            raise InterruptedError("Indexing cancelled")
        delete_file_from_index(path)
        meta_mtimes.pop(path, None)
        removed += 1
        step += 1
        set_progress(active=True, current=step, total=max(total_steps, 1), folder=folder, file=os.path.basename(path))

    updated = 0
    skipped = 0
    failed_paths = []
    pending = []

    def flush_pending():
        nonlocal updated
        if not pending:
            return
        paths = list(pending)
        pending.clear()
        try:
            counts = _index_text_batch(paths, model, max(1, int(os.environ.get("AETHER_EMBED_BATCH_SIZE", "16"))))
            for path in paths:
                mtime = os.path.getmtime(path)
                meta_mtimes[path] = (mtime, os.path.getsize(path))
                updated += 1
                _mark_progress()
        except InterruptedError:
            raise
        except Exception as batch_error:
            # A bad document must not discard otherwise valid work. Retry the
            # batch one file at a time, preserving the established failure and
            # replacement semantics for each file.
            print(f"INDEX BATCH FALLBACK: {batch_error}", file=sys.stderr)
            for path in paths:
                try:
                    count = _index_file(path, model=model)
                    if count:
                        updated += 1
                        meta_mtimes[path] = (os.path.getmtime(path), os.path.getsize(path))
                    else:
                        failed_paths.append(path)
                except (ExtractionError, OSError, ValueError, RuntimeError) as exc:
                    print(f"INDEX SKIP: {path}: {exc}", file=sys.stderr)
                    failed_paths.append(path)
                    meta_mtimes.pop(path, None)
                _mark_progress()

    def _mark_progress():
        nonlocal step
        step += 1
        set_progress(active=True, current=step, total=max(total_steps, 1), folder=folder, file="")

    for path in disk_paths:
        if is_cancelled():
            raise InterruptedError("Indexing cancelled")

        try:
            mtime = os.path.getmtime(path)
            size = os.path.getsize(path)
        except OSError:
            failed_paths.append(path)
            step += 1
            continue

        set_progress(
            active=True,
            current=step,
            total=max(total_steps, 1),
            folder=folder,
            file=os.path.basename(path),
        )

        previous_fingerprint = meta_mtimes.get(path)
        unchanged = (
            previous_fingerprint == mtime
            or previous_fingerprint == (mtime, size)
            or previous_fingerprint == [mtime, size]
        )
        if not force and path in indexed_paths and unchanged:
            skipped += 1
            step += 1
            continue
        if os.path.splitext(path)[1].lower() in TEXT_EXTENSIONS:
            pending.append(path)
            # Use a conservative byte-size estimate here so collection does
            # not re-extract every pending file just to decide when to flush.
            # `_index_text_batch` performs the authoritative chunking and
            # keeps model calls bounded even when one document is large.
            estimated_chunks = sum(
                max(1, math.ceil(os.path.getsize(p) / 650))
                for p in pending
                if os.path.exists(p)
            )
            if estimated_chunks >= max(1, int(os.environ.get("AETHER_EMBED_BATCH_SIZE", "16"))):
                flush_pending()
        else:
            flush_pending()
            try:
                chunk_count = _index_file(path, model=model)
                if chunk_count:
                    updated += 1
                    meta_mtimes[path] = (mtime, os.path.getsize(path))
                else:
                    failed_paths.append(path)
                    meta_mtimes.pop(path, None)
            except (ExtractionError, OSError, ValueError, RuntimeError) as exc:
                print(f"INDEX SKIP: {path}: {exc}", file=sys.stderr)
                failed_paths.append(path)
                meta_mtimes.pop(path, None)
            _mark_progress()

    flush_pending()

    for path in list(meta_mtimes):
        if path not in disk_path_set:
            del meta_mtimes[path]

    set_folder_file_mtimes(folder, meta_mtimes)
    stats = get_folder_stats(folder)
    failed_sample = failed_paths[:8]
    update_folder_meta(
        folder,
        file_count=stats["file_count"],
        chunk_count=stats["chunk_count"],
        skipped=skipped,
        last_index_failures=failed_sample,
    )
    return {
        "chunks": stats["chunk_count"],
        "updated": updated,
        "removed": removed,
        "skipped": skipped,
        "failed": len(failed_paths),
        "failed_sample": failed_sample,
    }


def build_index(folder, force=False):
    from index_progress import clear_progress

    try:
        return _build_index(folder, force=force)
    finally:
        # Never leave stale active/cancelled progress after success or failure.
        clear_progress()


def build_index_all(folders, force=False):
    from index_progress import clear_progress, is_cancelled

    totals = {
        "chunks": 0,
        "updated": 0,
        "removed": 0,
        "skipped": 0,
        "failed": 0,
        "failed_sample": [],
    }
    try:
        for folder in folders:
            if is_cancelled():
                raise InterruptedError("Indexing cancelled")
            folder = _normalize_folder(folder)
            if os.path.isdir(folder):
                result = _build_index(folder, force=force)
                for key in ("chunks", "updated", "removed", "skipped", "failed"):
                    totals[key] += result.get(key, 0)
                for path in result.get("failed_sample") or []:
                    if len(totals["failed_sample"]) >= 8:
                        break
                    if path not in totals["failed_sample"]:
                        totals["failed_sample"].append(path)
        return totals
    finally:
        clear_progress()


def index_single_file(path):
    path = os.path.normpath(path)
    chunk_count = _index_file(path)

    from folders import get_folder_file_mtimes, set_folder_file_mtimes

    folder = _parent_indexed_folder(path)
    if folder:
        mtimes = get_folder_file_mtimes(folder)
        if chunk_count:
            try:
                mtimes[path] = os.path.getmtime(path)
            except OSError:
                pass
        else:
            mtimes.pop(path, None)
        set_folder_file_mtimes(folder, mtimes)
