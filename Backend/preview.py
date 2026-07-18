import os

import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

from indexer import PDF_OCR_MIN_CHARS, TEXT_EXTENSIONS, ocr_pdf_pages


def preview_file(path):
    if not os.path.exists(path):
        return ""

    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        try:
            reader = PdfReader(path)
            text = ""
            for page in reader.pages[:3]:
                text += (page.extract_text() or "")
            if len(text.strip()) < PDF_OCR_MIN_CHARS:
                ocr_text = ocr_pdf_pages(path, max_pages=3)
                if len(ocr_text.strip()) > len(text.strip()):
                    return ocr_text
            return text
        except Exception:
            return ""

    elif ext == ".docx":
        try:
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return ""

    elif ext == ".xlsx":
        try:
            df = pd.read_excel(path)
            return df.to_string()
        except Exception:
            return ""

    elif ext == ".pptx":
        try:
            prs = Presentation(path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text
        except Exception:
            return ""

    elif ext in [".png", ".jpg", ".jpeg"]:
        try:
            # pyrefly: ignore [missing-import]
            from rapidocr_onnxruntime import RapidOCR
            engine = RapidOCR()
            result, _ = engine(path)
            if result:
                return "\n".join(line[1] for line in result if line and len(line) > 1)
            return ""
        except Exception:
            return ""

    elif ext in TEXT_EXTENSIONS:
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""

    # Fallback: attempt plain-text read for unknown extensions
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""
